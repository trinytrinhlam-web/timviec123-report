# -*- coding: utf-8 -*-
"""Trích dữ liệu nghiệm thu Giai đoạn 1 từ file QA 'Tester Detail- Timviec123'.

    python3 extract_nghiemthu_gd1.py <duong_dan_file_QA.xlsx> [ngay_chot dd/mm/yyyy]

Xuất data_nghiemthu_gd1.json để build_nghiemthu_gd1.py dựng biên bản.
Chỉ đếm là lỗi những dòng CÓ BUG_ID (dòng 'Không lỗi' không tính).
"""
import sys, json, re, datetime, collections
import openpyxl

TARGET=["DS_04","DS_05","DS_21","DS_22","DS_25"]
SCOPE={"DS_04":["YC_18","YC_19"],"DS_05":["YC_53"],
       "DS_21":["YC_01","YC_02","YC_03","YC_11"],
       "DS_22":["YC_04","YC_05","YC_07"],"DS_25":["YC_08","YC_09"]}

def rows(wb,name,hr):
    ws=wb[name]; hdr=[str(c.value).strip() if c.value is not None else "" for c in ws[hr]]
    out=[]
    for r in range(hr+1,ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1,len(hdr)+1)]
        if not any(v not in (None,"") for v in vals): continue
        out.append(dict(zip(hdr,vals)))
    return out

def d(x):
    if isinstance(x,datetime.datetime): return x.strftime("%d/%m/%Y")
    return "" if x is None else str(x).strip()

def main(path,asof):
    wb=openpyxl.load_workbook(path,data_only=True)
    ui={r["DS_ID"]:r for r in rows(wb,"02_UI_UX",4) if r.get("DS_ID")}
    tn=rows(wb,"03_Tính_năng",4)
    bugs=[b for b in rows(wb,"05_Lỗi_Tester",4) if b.get("BUG_ID")]
    rtm={}
    for r in rows(wb,"06_RTM_Dev_Test",4):
        if r.get("BUG_ID"): rtm.setdefault((r.get("DS_ID"),r["BUG_ID"]),r)
    tcs=rows(wb,"04_Test_Cases",4)

    out={"asof":asof,"ds":[]}
    for ds in TARGET:
        p=ui[ds]
        bs=[b for b in bugs if b.get("DS_ID")==ds]
        blist=[]
        for b in bs:
            k=rtm.get((ds,b["BUG_ID"]),{})
            blist.append({"bug":d(b.get("BUG_ID")),"yc":d(b.get("YC_ID")),"sev":d(b.get("Mức độ")),
                "type":d(b.get("Loại lỗi")),"dev_env":d(b.get("Thiết bị")),
                "title":d(b.get("Tiêu đề lỗi")),"expect":d(b.get("Kết quả mong đợi")),
                "found":d(b.get("Ngày phát hiện")),"dev":d(k.get("Dev phụ trách")),
                "fixed":d(k.get("Ngày Dev báo fix")),"retest":d(k.get("Kết quả Retest")),
                "retest_date":d(k.get("Ngày Retest")),"status":d(b.get("Trạng thái lỗi hiện tại")),
                "evi":bool(d(b.get("Bằng chứng thực tế (link)")) or d(k.get("Bằng chứng thực tế (link)"))),
                "fig":bool(d(b.get("Ảnh thiết kế (link)")))})
        ycs=[]
        for y in SCOPE[ds]:
            row=next((r for r in tn if d(r.get("REQ_ID"))==y and d(r.get("DS_ID"))==ds),None) \
                or next((r for r in tn if d(r.get("REQ_ID"))==y),{})
            nb=[b for b in bs if y in re.findall(r'YC_\d+',d(b.get("YC_ID")))]
            ycs.append({"yc":y,"name":d(row.get("Tên tính năng")),"prio":d(row.get("Ưu tiên")),
                "stat":d(row.get("Trạng thái kiểm tra")),"bugs_in_ds":len(nb),
                "open_in_ds":sum(1 for b in nb if d(b.get("Trạng thái lỗi hiện tại")) in ("Open","Chưa gửi Dev","Đang sửa")),
                "last":d(row.get("Ngày test gần nhất"))})
        rel=[t for t in tcs if ds in re.findall(r'DS_\d+',d(t.get("DS_ID(s)")))]
        out["ds"].append({"ds":ds,"name":d(p.get("Tên trang")),"url":d(p.get("Link Trang")),
            "stat":d(p.get("Trạng thái kiểm tra")),
            "dev":{k:d(p.get(k)) for k in ("Desktop","Mobile","Tablet")},
            "pct":p.get("Hoàn thành %"),"last":d(p.get("Ngày test gần nhất")),"tester":d(p.get("Tester")),
            "yc":ycs,"bugs":blist,"ntc":len(rel),
            "tcst":dict(collections.Counter(d(t.get("Trạng thái")) for t in rel))})
    json.dump(out,open("data_nghiemthu_gd1.json","w"),ensure_ascii=False,indent=1)
    for x in out["ds"]:
        op=sum(1 for b in x["bugs"] if b["status"] in ("Open","Chưa gửi Dev","Đang sửa"))
        print(f'{x["ds"]:7}{x["name"]:24} lỗi={len(x["bugs"]):2} mở={op} '
              f'retestPass={sum(1 for b in x["bugs"] if b["retest"]=="Pass")}')

if __name__=="__main__":
    main(sys.argv[1] if len(sys.argv)>1 else "/tmp/qa.xlsx",
         sys.argv[2] if len(sys.argv)>2 else datetime.date.today().strftime("%d/%m/%Y"))
