# -*- coding: utf-8 -*-
"""Biên bản nghiệm thu GĐ2 (Nhà tuyển dụng) — lần 2.

Đọc data_nghiemthu_gd2.json (do extract_nghiemthu_gd2.py sinh) và dựng
BaoCao_NghiemThu_GD2_lan2.xlsx. Bố cục giữ theo biên bản lần 1 (29/09/2026),
thêm cột đối chiếu Lần 1 -> Lần 2 ở mọi bảng.

Đánh giá từng mã lỗi ở lần 2:
  ĐÃ ĐẠT        kiểm lại trên website đạt, hoặc RTM có Retest Pass sau 29/09
  KHÔNG ĐẠT     kiểm lại trên website, lỗi vẫn còn
  CHỜ KIỂM LẠI  Dev báo đã sửa (Fixed/Verified) nhưng chưa có retest đạt
  CÒN MỞ        QA vẫn Open/Chưa gửi Dev/Đang sửa, chưa kiểm lại được
Kết luận ĐẠT chỉ khi mọi mã thuộc nhóm chặn duyệt đều ĐÃ ĐẠT.
"""
import json, os, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY = "1F3A5F"; BLUE = "2E5A88"; LBLUE = "DCE6F1"; ZEBRA = "F5F7FA"
GREEN = "2E7D32"; GREEN_BG = "D9F0DD"; RED = "C62828"; RED_BG = "FBE3E3"
AMBER = "8A6100"; AMBER_BG = "FFF4D6"; GREY = "6B7280"; WHITE = "FFFFFF"; INPUT = "FFF9DB"
thin = Side(style="thin", color="C9D6E5"); box = Border(left=thin, right=thin, top=thin, bottom=thin)


def F(size=10, bold=False, color="1A1A1A", italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def Fill(c): return PatternFill("solid", fgColor=c)


def cao(texts_widths, line_h=13.5, min_h=18):
    """Chiều cao dòng đủ chứa chữ: texts_widths = [(chuỗi, độ rộng cột gộp)]."""
    n = 1
    for txt, w in texts_widths:
        if txt:
            n = max(n, sum(-(-len(part) // max(1, int(w * 1.05))) or 1 for part in str(txt).split("\n")))
    return max(min_h, n * line_h + 4)
def Al(h="left", v="center", wrap=True): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


DG_STYLE = {"ĐÃ ĐẠT": (GREEN, GREEN_BG), "KHÔNG ĐẠT": (RED, RED_BG), "CÒN MỞ": (RED, RED_BG),
            "CHỜ KIỂM LẠI": (AMBER, AMBER_BG), "CLOSED": (GREY, ZEBRA), "DEFERRED": (GREY, ZEBRA),
            "Pass": (GREEN, GREEN_BG), "Fail": (RED, RED_BG), "Không xác định": (AMBER, AMBER_BG),
            "Đạt": (GREEN, GREEN_BG), "Đạt, có ghi chú": (GREEN, GREEN_BG), "Chưa đạt": (RED, RED_BG),
            "Nghiệm thu sau": (GREY, ZEBRA), "Ngoài phạm vi đợt này": (GREY, ZEBRA)}

data = json.load(open(os.path.join(HERE, "data_nghiemthu_gd2.json"), encoding="utf-8"))
ASOF = data["asof"]; META = data["kiem_lai_meta"]
items = data["items"]; by_ma = {i["ma"]: i for i in items}


def kl(it):  # kết quả kiểm lại 16/10 của một mã (dict hoặc {})
    return it.get("kiem_lai") or {}


def ma_tt_text(st):
    return "; ".join(f"{c}: {v}" for c, v in st.items())


# ---------------- tính kết luận ----------------
chan = [i for i in items if i["chan"]]
chan_con = [i for i in chan if i["danh_gia"] != "ĐÃ ĐẠT"]
def ket_luan_2(kl1, ma_tt, bad):
    """Kết luận lần 2 của luồng / yêu cầu / trang: còn mã chặn -> Chưa đạt; hết mã chặn thì
    giữ kết luận lần 1, riêng lần 1 'Chưa đạt' chuyển thành Đạt (có ghi chú nếu còn mã chưa đạt)."""
    if kl1 in ("Nghiệm thu sau", "Ngoài phạm vi đợt này"):
        return kl1
    if bad:
        return "Chưa đạt"
    if kl1 == "Chưa đạt":
        return "Đạt, có ghi chú" if any(v != "ĐÃ ĐẠT" for v in ma_tt.values()) else "Đạt"
    return kl1


for f in data["luong"]:
    f["kl2"] = ket_luan_2(f["kl1"], f["ma_tt"], f["chan_con"])
for y in data["yc"] + data["trang"]:
    y["chan_con"] = [c for c, v in y["ma_tt"].items() if by_ma[c]["chan"] and v != "ĐÃ ĐẠT"]
    y["kl2"] = ket_luan_2(y["kl1"], y["ma_tt"], y["chan_con"])
for k in data["kich_ban"]:
    st = k["ma_tt"]; kk = k.get("kiem_lai")
    if kk:
        k["kq2"] = {"Pass": "Đạt", "Fail": "Lỗi"}.get(kk["ket_qua"], "Chưa kiểm lại")
        k["gc2"] = kk.get("nhan_xet", "")
    elif not st:
        k["kq2"] = "Giữ kết quả 29/09" if k["kq1"] == "Đạt" else "Chưa kiểm lại"
        k["gc2"] = "Không có mã lỗi; chưa chạy lại ở lần 2" if k["kq1"] == "Đạt" else ""
    elif all(v == "ĐÃ ĐẠT" for v in st.values()):
        k["kq2"] = "Đạt"; k["gc2"] = "Các mã lỗi của kịch bản đã đạt khi kiểm lại"
    elif any(v == "KHÔNG ĐẠT" for v in st.values()):
        k["kq2"] = "Lỗi"; k["gc2"] = ma_tt_text(st)
    else:
        k["kq2"] = "Chưa kiểm lại"; k["gc2"] = ma_tt_text(st)

VERDICT = "ĐẠT" if not chan_con else "CHƯA ĐẠT"
cnt = collections.Counter
nhom_order = ["Cần sửa trước khi duyệt", "QA - cần xử lý", "Phát sinh - chặn duyệt", "Không chặn duyệt (tối ưu)",
              "Yêu cầu bổ sung", "QA - lệch nhỏ", "Phát sinh - không chặn", "Phát sinh - toàn hệ thống"]
NHOM_CHAN_ALL = ("Cần sửa trước khi duyệt", "QA - cần xử lý", "Phát sinh - chặn duyệt")

# ---------------- workbook ----------------
wb = openpyxl.Workbook()


def table(ws, headers, widths, rows, styles=None, hdr_row=1, freeze=True, heights=None):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(hdr_row, i, h); c.font = F(9.5, True, WHITE); c.fill = Fill(NAVY); c.alignment = Al("center"); c.border = box
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hdr_row].height = 32
    for r, row in enumerate(rows, hdr_row + 1):
        bg = Fill(ZEBRA) if (r - hdr_row) % 2 == 0 else Fill(WHITE)
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v if v != "" else None); c.font = F(9.5); c.fill = bg; c.border = box
            c.alignment = Al("left" if widths[i - 1] >= 18 else "center", "top")
            if styles and v in DG_STYLE and i - 1 in styles:
                fc, bc = DG_STYLE[v]; c.font = F(9.5, True, fc); c.fill = Fill(bc); c.alignment = Al("center")
        if heights:
            ws.row_dimensions[r].height = heights
    if freeze:
        ws.freeze_panes = ws.cell(hdr_row + 1, 3)
    ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(len(headers))}{hdr_row + len(rows)}"
    ws.sheet_view.showGridLines = False


# ===== Sheet 1: Biên bản =====
ws = wb.active; ws.title = "Biên bản"
for col, w in {"A": 2, "B": 40, "C": 20, "D": 20, "E": 20, "F": 46}.items():
    ws.column_dimensions[col].width = w
ws.sheet_view.showGridLines = False
row = [2]


def line(label, val, lab_font=None, val_font=None, val_fill=None, h=None, merge_to="F"):
    r = row[0]
    c = ws.cell(r, 2, label); c.font = lab_font or F(10, True); c.alignment = Al("left", "top"); c.border = box
    ws.merge_cells(f"C{r}:{merge_to}{r}")
    c = ws.cell(r, 3, val); c.font = val_font or F(10); c.alignment = Al("left", "top")
    for col in range(3, 7):
        ws.cell(r, col).border = box
        if val_fill: ws.cell(r, col).fill = Fill(val_fill)
    ws.row_dimensions[r].height = max(h or 0, cao([(label, 40), (val, 106)]))
    row[0] += 1


def head(text):
    r = row[0]; ws.merge_cells(f"B{r}:F{r}")
    c = ws.cell(r, 2, text); c.font = F(11.5, True, WHITE); c.fill = Fill(BLUE); c.alignment = Al("left")
    for col in range(2, 7): ws.cell(r, col).fill = Fill(BLUE)
    ws.row_dimensions[r].height = 22; row[0] += 1


def gap(): row[0] += 1


for txt, fnt, h in [("BIÊN BẢN NGHIỆM THU — LẦN 2", F(16, True, NAVY), 28),
                    ("Website timviec123.vn — Đợt 2: Nhà tuyển dụng (V3)", F(12, True, BLUE), 20),
                    (f"Ngày lập: {ASOF}  ·  Kiểm lại: {META.get('ngay_kiem', ASOF)}  ·  Bản giao diện: {META.get('ban_giao_dien', '')}"
                     f"  ·  Môi trường: website đang vận hành https://timviec123.vn  ·  Số liệu QA đến {ASOF}"
                     f"  ·  Lần 1: biên bản 29/09/2026 (CHƯA ĐẠT)", F(9.5, italic=True, color=GREY), 30)]:
    r = row[0]; ws.merge_cells(f"B{r}:F{r}"); c = ws.cell(r, 2, txt); c.font = fnt; c.alignment = Al("left")
    ws.row_dimensions[r].height = h; row[0] += 1
gap()

head("1. PHẠM VI")
line("Tác nhân", "Nhà tuyển dụng (NTD), kèm phía Ứng viên và Khách ở các luồng liên tác — giữ nguyên phạm vi lần 1.")
line("Nguồn nghiệm thu", "(1) 38 mã lỗi / yêu cầu của biên bản lần 1 và các lỗi phát sinh sau 29/09 trong khu NTD; "
     "(2) 116 kịch bản, 8 luồng F1–F8, 23 yêu cầu tính năng, 19 trang NTD của lần 1; "
     f"(3) {len(data['verified'])} lỗi đang ở trạng thái Verified trong 06_RTM_Dev_Test (Dev báo đã sửa, chờ tester kiểm).", h=44)
line("Cách kiểm lần 2", META.get("cach_kiem", ""), h=58)
line("Giới hạn lần 2", META.get("gioi_han", ""), val_fill=AMBER_BG, h=58)
line("Chưa đưa vào", "Gói dịch vụ trả phí, Footer, Tạo ảnh chân dung, chức năng Admin (Đợt 3) — như lần 1. "
     "BUG_API_004 (làm mới token) ảnh hưởng mọi vai trò: ghi nhận, không tính vào điều kiện duyệt GĐ2.", h=30)
gap()

head("2. TIÊU CHUẨN DUYỆT")
line("Tính năng", "Hoạt động, người dùng thao tác không gặp lỗi")
line("Giao diện", "Đúng thiết kế Figma; lỗi làm hỏng bố cục hoặc chặn thao tác phải sửa để đạt; lệch nhỏ chuyển tối ưu", h=28)
line("Yêu cầu bổ sung", "Mã REQ là yêu cầu chưa có trên website, không tính là lỗi; theo dõi riêng")
line("Điều kiện đạt lần 2", f"Toàn bộ {len(chan)} mã thuộc nhóm chặn duyệt phải ĐÃ ĐẠT khi kiểm lại: 8 lỗi 'Cần sửa trước khi duyệt' "
     f"+ 7 lỗi 'QA - cần xử lý' của lần 1 + {sum(1 for i in chan if i['phat_sinh'])} lỗi Critical/High phát sinh trong khu NTD. "
     "Dev báo đã sửa nhưng chưa retest đạt thì chưa tính.", h=44)
gap()

head("3. KẾT QUẢ — ĐỐI CHIẾU LẦN 1 / LẦN 2")
r = row[0]
for col, txt in zip(range(2, 7), ["Chỉ số", "Lần 1 (29/09)", f"Lần 2 ({ASOF})", "", "Ghi chú"]):
    c = ws.cell(r, col, txt); c.font = F(10, True, WHITE); c.fill = Fill(NAVY); c.alignment = Al("center"); c.border = box
ws.merge_cells(f"D{r}:E{r}"); row[0] += 1


def kq(label, v1, v2, note="", bad=False):
    r = row[0]
    vals = [label, v1, v2, None, note]
    for col, v in zip(range(2, 7), vals):
        c = ws.cell(r, col, v); c.border = box; c.font = F(10); c.alignment = Al("left" if col in (2, 6) else "center")
    ws.merge_cells(f"D{r}:E{r}")
    if bad:
        ws.cell(r, 4).font = F(10, True, RED)
    ws.row_dimensions[r].height = cao([(label, 40), (v1, 20), (v2, 40), (note, 46)])
    row[0] += 1


def dg_cnt(lst):
    c = cnt(i["danh_gia"] for i in lst)
    return ", ".join(f"{k.lower()} {v}" for k, v in c.most_common())


for nh, lab in [("Cần sửa trước khi duyệt", "Lỗi cần sửa trước khi duyệt"), ("QA - cần xử lý", "Lỗi giao diện QA cần xử lý"),
                ("Phát sinh - chặn duyệt", "Lỗi C/H phát sinh sau 29/09 (khu NTD)"),
                ("Không chặn duyệt (tối ưu)", "Lỗi không chặn duyệt"), ("Yêu cầu bổ sung", "Yêu cầu bổ sung (REQ)"),
                ("QA - lệch nhỏ", "Lệch thiết kế nhỏ"), ("Phát sinh - không chặn", "Lỗi khác phát sinh / mở lại (khu NTD)")]:
    lst = [i for i in items if i["nhom"] == nh]
    if not lst:
        continue
    da = sum(1 for i in lst if i["danh_gia"] == "ĐÃ ĐẠT")
    v1 = "—" if nh.startswith("Phát sinh") else f"{len(lst)} mở"
    kq(lab, v1, f"{da}/{len(lst)} đạt", dg_cnt(lst), bad=da < len(lst) and nh in NHOM_CHAN_ALL)
fl2 = cnt(f["kl2"] for f in data["luong"]); fl1 = cnt(f["kl1"] for f in data["luong"])
kq("Luồng F1–F8", ", ".join(f"{k}: {v}" for k, v in fl1.items()), ", ".join(f"{k}: {v}" for k, v in fl2.items()),
   "Luồng chưa đạt khi còn mã chặn duyệt của luồng hoặc của trang mà luồng đi qua")
y1 = cnt(y["kl1"] for y in data["yc"]); y2 = cnt(y["kl2"] for y in data["yc"])
kq("Yêu cầu tính năng (23)", ", ".join(f"{k}: {v}" for k, v in y1.items()), ", ".join(f"{k}: {v}" for k, v in y2.items()))
k1 = cnt(k["kq1"] for k in data["kich_ban"]); k2 = cnt(k["kq2"] for k in data["kich_ban"])
kq("Kịch bản (116)", ", ".join(f"{k}: {v}" for k, v in k1.items()), ", ".join(f"{k}: {v}" for k, v in k2.items()),
   "'Giữ kết quả 29/09' = kịch bản đã Đạt ở lần 1, không có mã lỗi, chưa chạy lại")
vv = data["verified"]; vk = cnt((v["kiem_lai"] or {}).get("ket_qua", "Chưa kiểm") for v in vv)
kq(f"Lỗi Verified toàn site ({len(vv)})", "—", ", ".join(f"{k}: {v}" for k, v in vk.items()),
   "Sheet 'Lỗi Verified (Dev báo sửa)'. Chỉ ghi kết quả ở báo cáo, chưa ghi vào file QA")
gap()

head("4. KẾT LUẬN")
r = row[0]; ws.merge_cells(f"B{r}:F{r}")
c = ws.cell(r, 2, VERDICT); c.font = F(16, True, GREEN if VERDICT == "ĐẠT" else RED)
c.fill = Fill(GREEN_BG if VERDICT == "ĐẠT" else RED_BG); c.alignment = Al("center")
for col in range(2, 7): ws.cell(r, col).fill = Fill(GREEN_BG if VERDICT == "ĐẠT" else RED_BG)
ws.row_dimensions[r].height = 30; row[0] += 1
if chan_con:
    g = collections.defaultdict(list)
    for i in chan_con: g[i["danh_gia"]].append(i["ma"])
    ly_do = (f"Còn {len(chan_con)}/{len(chan)} mã chặn duyệt chưa đạt — "
             + "; ".join(f"{k} ({len(v)}): {', '.join(v)}" for k, v in g.items()) + ". " + META.get("ket_luan_them", ""))
else:
    ly_do = f"Toàn bộ {len(chan)} mã chặn duyệt đã đạt khi kiểm lại. " + META.get("ket_luan_them", "")
line("Lý do", ly_do)
gap()

head("5. ĐỀ NGHỊ")
for t in META.get("de_nghi", []):
    r = row[0]; ws.merge_cells(f"B{r}:F{r}")
    c = ws.cell(r, 2, "•  " + t); c.font = F(10); c.alignment = Al("left", "top")
    ws.row_dimensions[r].height = 15 * (len(t) // 110 + 1) + 4; row[0] += 1
gap()

head("6. XÁC NHẬN")
r = row[0]
for col, t in [(2, "Người lập (QA)"), (4, "Đại diện phát triển"), (6, "Ban Giám đốc")]:
    ws.cell(r, col, t).font = F(10, True); ws.cell(r + 1, col, "(Ký, ghi rõ họ tên)").font = F(9, italic=True, color=GREY)
    ws.cell(r, col).alignment = Al("center"); ws.cell(r + 1, col).alignment = Al("center")
row[0] += 6

# ===== Sheet 2: Điểm chặn =====
HDR = ["STT", "Mã", "Nhóm", "Mức độ", "DS_ID", "Trang", "Tiêu đề", "Thiết bị", "Lần 1 (29/09)", "Trạng thái QA",
       "Dev báo fix", "Kiểm lại", "Nhận xét kiểm lại", "Bằng chứng mã nguồn (bản giao diện mới)", "ĐÁNH GIÁ LẦN 2"]
WID = [5, 17, 18, 9, 10, 22, 40, 14, 22, 12, 11, 12, 48, 40, 15]


def item_row(n, i):
    k = kl(i)
    return [n, i["ma"], i["nhom"], i["muc_do"], i["ds"], i["trang"], i["tieu_de"], i["thiet_bi"], i["lan1"], i["tt_qa"],
            i["ngay_fix"] or "—", k.get("ket_qua", "Chưa kiểm"), k.get("nhan_xet", ""), k.get("bang_chung_ma", ""), i["danh_gia"]]


ws2 = wb.create_sheet("Điểm chặn")
srt = sorted(chan, key=lambda i: (nhom_order.index(i["nhom"]), i["ma"]))
table(ws2, HDR, WID, [item_row(n, i) for n, i in enumerate(srt, 1)], styles={11, 14}, heights=60)

ws3 = wb.create_sheet("Danh sách lỗi lần 2")
srt = sorted(items, key=lambda i: (nhom_order.index(i["nhom"]), i["ma"]))
table(ws3, HDR, WID, [item_row(n, i) for n, i in enumerate(srt, 1)], styles={11, 14}, heights=60)

# ===== Luồng / YC / Trang / Kịch bản =====
ws4 = wb.create_sheet("Kết quả theo luồng")
table(ws4, ["Luồng", "Tên luồng", "Số kịch bản", "Đạt / Lỗi / Ghi nhận (29/09)", "Kết luận lần 1",
            "Mã lỗi của luồng và trang đi qua — đánh giá lần 2", "Mã chặn duyệt còn lại", "Kết luận lần 2"],
      [7, 30, 9, 14, 15, 70, 40, 15],
      [[f["ma"], f["ten"], f["so_kb"], f"{f['dat']} / {f['loi']} / {f['ghi_nhan']}", f["kl1"], ma_tt_text(f["ma_tt"]),
        ", ".join(f["chan_con"]) or "—", f["kl2"]] for f in data["luong"]], styles={4, 7}, heights=75)

ws5 = wb.create_sheet("Tính năng theo yêu cầu")
table(ws5, ["Mã yêu cầu", "Tên yêu cầu", "Trang", "Ưu tiên", "Kết luận lần 1", "Mã lỗi liên quan — đánh giá lần 2",
            "Mã chặn còn lại", "Kết luận lần 2"],
      [10, 32, 14, 8, 15, 60, 30, 15],
      [[y["ma"], y["ten"], y["trang"], y["uu_tien"], y["kl1"], ma_tt_text(y["ma_tt"]) or "—",
        ", ".join(y["chan_con"]) or "—", y["kl2"]] for y in data["yc"]], styles={4, 7}, heights=45)

ws6 = wb.create_sheet("Giao diện theo trang")
table(ws6, ["DS_ID", "Tên trang", "Tổng lỗi QA (29/09)", "Còn mở (29/09)", "Lỗi mở trong QA hiện tại", "Kết luận lần 1",
            "Mã lỗi — đánh giá lần 2", "Mã chặn còn lại", "Kết luận lần 2"],
      [8, 32, 10, 10, 11, 15, 60, 30, 15],
      [[p["ds"], p["ten"], p["tong"], p["con_mo"], p["tt_qa_moi"], p["kl1"], ma_tt_text(p["ma_tt"]) or "—",
        ", ".join(p["chan_con"]) or "—", p["kl2"]] for p in data["trang"]], styles={5, 8}, heights=45)

ws7 = wb.create_sheet("Kịch bản chi tiết")
table(ws7, ["Mã", "Luồng", "Kịch bản", "Ưu tiên", "Thiết bị", "Trang", "YC", "Kết quả 29/09", "Kết luận 29/09",
            "Mã lỗi / yêu cầu", "Kết quả lần 2", "Ghi chú lần 2"],
      [11, 22, 34, 7, 8, 10, 9, 10, 15, 26, 16, 50],
      [[k["ma"], k["luong"], k["ten"], k["uu_tien"], k["thiet_bi"], k["ds"], k["yc"], k["kq1"], k["kl1"], k["ma_loi"],
        k["kq2"], k["gc2"]] for k in data["kich_ban"]], heights=40)
for r in range(2, ws7.max_row + 1):
    v = ws7.cell(r, 11).value
    col = {"Đạt": (GREEN, GREEN_BG), "Lỗi": (RED, RED_BG), "Chưa kiểm lại": (AMBER, AMBER_BG)}.get(v)
    if col:
        ws7.cell(r, 11).font = F(9.5, True, col[0]); ws7.cell(r, 11).fill = Fill(col[1])

# ===== Lỗi Verified =====
ws8 = wb.create_sheet("Lỗi Verified (Dev báo sửa)")


def de_xuat(v):
    k = (v["kiem_lai"] or {}).get("ket_qua")
    if k == "Pass": return "Ghi Retest = Pass, ngày " + META.get("ngay_kiem", ASOF) + "; chuyển Closed"
    if k == "Fail": return "Ghi Retest = Fail, ngày " + META.get("ngay_kiem", ASOF) + "; trả về Open cho Dev"
    return "Giữ Verified, chờ kiểm (cần tài khoản đã kích hoạt)" if not k or k == "Không xác định" else ""


vrows = []
for n, v in enumerate(sorted(data["verified"], key=lambda x: (not x["ntd"], x["ds"], x["ma"])), 1):
    k = v["kiem_lai"] or {}
    vrows.append([n, v["ma"], "Có" if v["ntd"] else "", v["ds"], v["trang"], v["muc_do"], v["thiet_bi"], v["tieu_de"],
                  f"{v['retest_rtm'] or '—'} {v['ngay_retest_rtm']}".strip(), k.get("ket_qua", "Chưa kiểm"),
                  k.get("nhan_xet", "") or ("Cần đăng nhập Ứng viên/NTD đã kích hoạt" if not k else ""),
                  k.get("bang_chung_ma", ""), de_xuat(v)])
table(ws8, ["STT", "Mã", "Khu NTD (GĐ2)", "DS_ID", "Trang", "Mức độ", "Thiết bị", "Tiêu đề", "Retest trước đó (RTM)",
            "Kiểm 16/10", "Nhận xét", "Bằng chứng mã nguồn", "Đề xuất cập nhật QA"],
      [5, 18, 8, 8, 22, 9, 14, 34, 14, 12, 50, 34, 30], vrows, styles={9}, heights=50)

# ===== Phát hiện mới =====
moi = META.get("phat_hien_moi", [])
if moi:
    wsm = wb.create_sheet("Phát hiện mới 16-10")
    table(wsm, ["STT", "Mô tả", "Mức đề xuất", "Trang", "Ảnh"], [5, 90, 14, 14, 40],
          [[n, m[0], m[1], m[2], ", ".join(os.path.basename(a) for a in m[3])] for n, m in enumerate(moi, 1)],
          freeze=False, heights=45)

# ===== Ảnh =====
ws9 = wb.create_sheet("Ảnh kiểm lại")
ws9.column_dimensions["A"].width = 2; ws9.column_dimensions["B"].width = 110; ws9.sheet_view.showGridLines = False
c = ws9.cell(1, 2, f"Ảnh chụp khi kiểm lại ngày {META.get('ngay_kiem', ASOF)} trên website thật (Chromium, khổ ghi ở từng ảnh)")
c.font = F(11, True, NAVY)
r = 3
for i in sorted(items, key=lambda i: (nhom_order.index(i["nhom"]), i["ma"])) + [
        {"ma": v["ma"], "kiem_lai": v["kiem_lai"], "danh_gia": "", "tieu_de": v["tieu_de"]} for v in data["verified"]
        if v["ma"] not in by_ma] + [
        {"ma": f"Phát hiện mới {n}", "kiem_lai": {"ket_qua": m[1], "anh": m[3]}, "tieu_de": m[0]}
        for n, m in enumerate(moi, 1) if m[3]]:
    k = kl(i)
    for a in k.get("anh", []):
        p = os.path.join(HERE, a)
        if not os.path.exists(p):
            continue
        ws9.cell(r, 2, f"{i['ma']}  ·  {k.get('ket_qua', '')}  ·  {os.path.basename(a)}  ·  {i['tieu_de'][:90]}").font = F(10, True)
        img = XLImage(p)
        scale = min(1.0, 620 / img.width, 520 / img.height)
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws9.add_image(img, f"B{r + 1}")
        r += int(img.height / 20) + 3

# ===== Cách kiểm =====
ws10 = wb.create_sheet("Cách kiểm lần 2")
ws10.column_dimensions["A"].width = 26; ws10.column_dimensions["B"].width = 110; ws10.sheet_view.showGridLines = False
for n, (a, b) in enumerate(META.get("phuong_phap", []), 1):
    ws10.cell(n, 1, a).font = F(10, True); c = ws10.cell(n, 2, b); c.font = F(10); c.alignment = Al("left", "top")
    ws10.cell(n, 1).alignment = Al("left", "top"); ws10.row_dimensions[n].height = 15 * (len(b) // 120 + 1) + 4

for w_ in wb.worksheets:
    w_.page_setup.orientation = "landscape"; w_.page_setup.paperSize = w_.PAPERSIZE_A4
    w_.sheet_properties.pageSetUpPr.fitToPage = True; w_.page_setup.fitToWidth = 1; w_.page_setup.fitToHeight = 0
    w_.print_options.horizontalCentered = True
    w_.page_margins.left = w_.page_margins.right = 0.4
ws.page_setup.orientation = "portrait"
for w_, clr in [(ws, NAVY), (ws2, RED), (ws3, BLUE), (ws8, AMBER)]:
    w_.sheet_properties.tabColor = clr
out = os.path.join(HERE, "BaoCao_NghiemThu_GD2_lan2.xlsx")
wb.save(out)
print("SAVED", out)
print("KẾT LUẬN:", VERDICT, "| chặn còn lại:", len(chan_con), "/", len(chan))
print("Luồng:", dict(fl2), "| YC:", dict(y2), "| Kịch bản:", dict(k2))
